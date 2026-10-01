# -*- coding: utf-8 -*-
"""自检：模板含「填写示例」sheet，且导入解析读不到示例行。"""
import sys
from io import BytesIO

sys.path.insert(0, ".")

from openpyxl import load_workbook

from app.core.excel import read_xlsx
from app.services.import_service import ImportService

failures = []


def check(name, cond, detail=""):
    tag = "PASS" if cond else "FAIL"
    print(f"[{tag}] {name} {detail}")
    if not cond:
        failures.append(name)


for kind, svc, sheet in (
    ("words", ImportService.words_template, "英语词条"),
    ("texts", ImportService.texts_template, "语文课文"),
):
    content = svc()
    wb = load_workbook(BytesIO(content))
    check(f"{kind}: sheet 列表", "填写示例" in wb.sheetnames, f"sheets={wb.sheetnames}")
    check(f"{kind}: active 是数据表", wb.active.title == sheet, f"active={wb.active.title}")
    # 导入解析口径（read_xlsx 默认活动表）：空模板必须 0 行 —— 示例绝不泄漏
    rows = read_xlsx(content)
    check(f"{kind}: 空模板导入解析 0 行", rows == [], f"rows={len(rows)}")

# 抽查示例表内容
wb = load_workbook(BytesIO(ImportService.words_template()))
ex = wb["填写示例"]
texts = [str(c.value or "") for row in ex.iter_rows() for c in row if c.value]
check("words 示例含 hello 行", any("hello" in t for t in texts))
check("words 示例含可接受答案示例", any("喂,哈罗" in t for t in texts))
wb2 = load_workbook(BytesIO(ImportService.texts_template()))
ex2 = wb2["填写示例"]
t2 = [str(c.value or "") for row in ex2.iter_rows() for c in row if c.value]
check("texts 示例含背诵范围示例 0", any(t == "0" for t in t2))
check("texts 示例含两段正文(Alt+Enter 换行)", any("\n" in t for t in t2))
check("texts 示例说明提到从 0 数起", any("从 0 开始" in t for t in t2))

# 学生名单模板
from app.core.excel import build_template

content = build_template(
    ["学号", "姓名", "班级", "初始密码"],
    sheet_name="学生名单",
    example_notes=["· 必填：学号、姓名"],
    example_rows=[["2026001", "张小明", "三年级(1)班", ""]],
)
wb3 = load_workbook(BytesIO(content))
check("students: active 是数据表", wb3.active.title == "学生名单", f"active={wb3.active.title}")
check("students: 示例表存在", "填写示例" in wb3.sheetnames)

print()
print("FAILURES:", failures if failures else "none")
sys.exit(1 if failures else 0)
