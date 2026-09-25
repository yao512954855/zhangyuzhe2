"""WP-BE-07 统计看板与申诉复核 —— 真机烟测（真 uvicorn + 真 SQLite 文件 + 真 HTTP）。

与 `app/tests/test_be07_stats.py` 的分工：单测用 TestClient（同进程、同事件循环）
覆盖口径与边界；**这里只验「真进程 + 真文件 + 真 HTTP」这件事**：

- 统计 SQL 在**独立进程**打开的同一个 SQLite 文件上跑，索引/排序/时区修饰符都真实生效
  （TestClient 里一切都在一个连接池里，某些并发/锁行为是测不出来的）
- xlsx 导出真的能作为二进制通过 HTTP 传输并被 openpyxl 打开
- 复核的写入真的落盘（换个进程/换条连接还能读到）

跑法：
    cd backend
    ./.venv/Scripts/python.exe scripts/smoke_be07.py

覆盖：U1（班级概览）、U2（单元掌握度）、U3（个人报告）、U4（正确率排除 empty）、
U5（跨接口口径一致）、U6（申诉 pending）、U7（复核不改 is_correct + 审计）、
U8（audio_available=false）、U10（ASR 用量）、U12（审计查询）、U13（xlsx 导出）、
U11（北京时间切天）+ 鉴权。

⚠️ 需要 TCP 8895 空闲；本脚本**不占用** UDP 9527（但若本机已有后端在跑，
   它的调度器可能同时动这个库 —— 一般无碍，只是别同时开着线上服务跑）。
"""

from __future__ import annotations

import io
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
PYTHON = BACKEND / ".venv" / "Scripts" / "python.exe"
if not PYTHON.exists():      # pragma: no cover - 非 Windows / 未建 venv
    PYTHON = Path(sys.executable)

PORT = int(os.environ.get("SMOKE_BE07_PORT", "8895"))
BASE = f"http://127.0.0.1:{PORT}"

ADMIN_USER = "admin"
ADMIN_PASS = "Smoke-Be07-Pass"
STUDENT_PASS = "Stu-Pass-BE07"

PASS = 0
FAIL = 0
FAILED: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"[PASS] {name}" + (f"  |  {detail}" if detail else ""), flush=True)
    else:
        FAIL += 1
        FAILED.append(name)
        print(f"[FAIL] {name}  |  {detail}", flush=True)


TMP = Path(tempfile.mkdtemp(prefix="smoke_be07_"))
DB_FILE = TMP / "smoke.db"

ENV = {
    **os.environ,
    "APP_PORT": str(PORT),
    "DB_PATH": str(DB_FILE),
    "BACKUP_DIR": str(TMP / "backups"),
    "JWT_SECRET": "smoke-be07-secret-0123456789abcdef",
    "JWT_TTL_HOURS": "12",
    "ADMIN_INIT_USERNAME": ADMIN_USER,
    "ADMIN_INIT_PASSWORD": ADMIN_PASS,
    "LOG_LEVEL": "WARNING",
    "NO_PROXY": "127.0.0.1,localhost",
    "PYTHONIOENCODING": "utf-8",
    # ⚠️ 造数脚本写在**临时目录**里，`sys.path[0]` 于是是那个临时目录，
    # 光靠 `cwd=backend` 是 import 不到 `app` 的（Python 不把 cwd 放进 sys.path）。
    # 其它烟测（be02/be03/be04）是脚本内 `sys.path.insert(0, BACKEND)`；
    # 这里子进程改不了它的 sys.path，只能靠 PYTHONPATH 传进去。
    "PYTHONPATH": str(BACKEND),
}


# ---------------------------------------------------------------------------
# 造数脚本：跑在**独立进程**里，用应用自己的 DB 层写库
# ---------------------------------------------------------------------------

SEED_SCRIPT = r'''
import json
from datetime import datetime, timedelta, timezone

from app.core.db import write_tx
from app.main import bootstrap

bootstrap()


def utc(bj_day, hour=12, minute=0):
    """北京时间 → UTC ISO 串（入库用）。"""
    dt = datetime.fromisoformat(f"{bj_day}T{hour:02d}:{minute:02d}:00+08:00")
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


CST = timezone(timedelta(hours=8))
today = datetime.now(timezone.utc).astimezone(CST).date().isoformat()
pattern = ["exact", "exact", "wrong", "empty"]      # 2 对 1 错 1 空 → 正确率 2/3

with write_tx() as conn:
    cid = conn.execute(
        "INSERT INTO classes (name, grade, semester, school_year) VALUES (?,?,?,?)",
        ("烟测五年级(1)班", 5, "first", "2026-2027"),
    ).lastrowid

    uid = conn.execute(
        "INSERT INTO units (subject, grade, semester, name, order_no) VALUES (?,?,?,?,?)",
        ("english", 5, "first", "烟测 Unit 1", 1),
    ).lastrowid

    word_ids = []
    for i in range(4):
        word_ids.append(conn.execute(
            "INSERT INTO words (unit_id, type, en, cn, order_no) VALUES (?,?,?,?,?)",
            (uid, "word", f"smoke{i}", f"烟测词{i}", i),
        ).lastrowid)

    student_ids = []
    for i in range(5):
        student_ids.append(conn.execute(
            """INSERT INTO users (role, username, real_name, password_hash,
                                  class_id, register_source, status)
               VALUES ('student', ?, ?, 'x', ?, 'import', 'active')""",
            (f"smoke{i:03d}", f"烟测学生{i}", cid),
        ).lastrowid)

    appealed = None
    for sid in student_ids:
        session_id = conn.execute(
            """INSERT INTO sessions (student_id, unit_id, content_type, strategy,
                                     total_items, correct_items, status, started_at, ended_at)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (sid, uid, "english", "sequential", 4, 2, "finished",
             utc(today, 10, 0), utc(today, 10, 3)),
        ).lastrowid

        for wi, mt in enumerate(pattern):
            answer_id = conn.execute(
                """INSERT INTO answers (session_id, student_id, word_id, mode, step,
                                        expected, asr_text, is_correct, match_type,
                                        duration_ms, created_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (session_id, sid, word_ids[wi], "cn2en", 1, f"smoke{wi}",
                 "" if mt == "empty" else "heard",
                 1 if mt in ("exact", "loose") else 0, mt, 3000, utc(today, 10, wi)),
            ).lastrowid
            if mt == "empty":
                continue        # BE-04 对 empty 不建掌握度行
            conn.execute(
                """INSERT INTO word_mastery (student_id, word_id, streak, total_count,
                                            correct_count, wrong_count, last_result,
                                            last_answer_at, mastered_at, avg_duration_ms)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (sid, word_ids[wi], 1, 1,
                 1 if mt in ("exact", "loose") else 0,
                 0 if mt in ("exact", "loose") else 1,
                 1 if mt in ("exact", "loose") else 0,
                 utc(today, 10, wi), utc(today, 10, wi) if wi == 0 else None, 3000),
            )
            if mt == "wrong" and appealed is None:
                conn.execute("UPDATE answers SET appealed = 1 WHERE id = ?", (answer_id,))
                appealed = answer_id

    # ASR 明细：4 次 aliyun，其中 1 次 error → 失败率 0.25
    for i in range(3):
        conn.execute(
            """INSERT INTO asr_logs (provider, model, lang, task_type, audio_ms,
                                     latency_ms, status, created_at)
               VALUES ('aliyun', 'fun-asr-flash', 'en', 'word', 1200, 250, 'ok', ?)""",
            (utc(today, 11, i),),
        )
    conn.execute(
        """INSERT INTO asr_logs (provider, model, lang, task_type, audio_ms,
                                 latency_ms, status, created_at)
           VALUES ('aliyun', 'fun-asr-flash', 'en', 'word', 900, 500, 'error', ?)""",
        (utc(today, 11, 30),),
    )

print(json.dumps({
    "class_id": cid, "unit_id": uid, "word_ids": word_ids,
    "student_ids": student_ids, "appealed_answer_id": appealed, "today": today,
}))
'''


def _tcp_free(port: int) -> bool:
    with socket.socket() as s:
        return s.connect_ex(("127.0.0.1", port)) != 0


def _http(method: str, path: str, *, token: str | None = None,
          json_body: dict | None = None, timeout: float = 30.0, raw: bool = False):
    import httpx

    headers = {"Authorization": f"Bearer {token}"} if token else {}
    with httpx.Client(timeout=timeout, trust_env=False) as c:
        resp = c.request(method, BASE + path, json=json_body, headers=headers)
    if raw:
        return resp.status_code, resp.content, dict(resp.headers)
    try:
        return resp.status_code, resp.json()
    except Exception:      # noqa: BLE001
        return resp.status_code, {"raw": resp.text[:200]}


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------


def run_checks(seed: dict) -> None:
    cid = seed["class_id"]
    uid = seed["unit_id"]
    today = seed["today"]

    # ---- 登录 ----
    status, body = _http("POST", "/api/auth/login",
                         json_body={"username": ADMIN_USER, "password": ADMIN_PASS,
                                    "role": "admin"})
    check("管理者登录", body.get("code") == 0, f"{str(body)[:140]}")
    admin_token = (body.get("data") or {}).get("access_token")
    if not admin_token:
        check("拿到管理员 token", False, "后续无法继续")
        return

    status, body = _http("POST", "/api/auth/register",
                         json_body={"username": "smokeout", "password": STUDENT_PASS,
                                    "real_name": "烟测外部学生", "grade": 5, "semester": 1})
    if body.get("code") != 0:
        status, body = _http("POST", "/api/auth/login",
                             json_body={"username": "smokeout", "password": STUDENT_PASS,
                                        "role": "student"})
    student_token = (body.get("data") or {}).get("access_token")
    check("学生登录（用于鉴权反例）", bool(student_token), f"{str(body)[:120]}")

    # ---- U1 班级概览 ----
    status, body = _http("GET", f"/api/stats/class/{cid}", token=admin_token)
    check("U1 班级概览 code=0", body.get("code") == 0, f"{str(body)[:160]}")
    d = body.get("data") or {}
    check("U1 学生数 = 5", d.get("student_count") == 5, f"student_count={d.get('student_count')}")
    check("U1 今日活跃 = 5（北京时间切天生效）",
          d.get("today_active_students") == 5 and d.get("today_date") == today,
          f"active={d.get('today_active_students')} today={d.get('today_date')}")
    check("U1 趋势 = 7 个点（缺天补零）", len(d.get("trend") or []) == 7,
          f"n={len(d.get('trend') or [])}")
    # 5 人 × (2 exact + 1 wrong)，empty 不进分母 → 10/15
    check("U1 正确率 = 10/15 ≈ 0.6667", d.get("accuracy") == 0.6667,
          f"accuracy={d.get('accuracy')} counted={d.get('counted_count')}")
    check("U1 分母排除 empty（15 而不是 20）",
          d.get("counted_count") == 15 and d.get("answer_count") == 20,
          f"counted={d.get('counted_count')} answer={d.get('answer_count')}")
    check("U1 错题 Top 有值且次数正确",
          bool(d.get("top_wrong_words")) and d["top_wrong_words"][0]["error_count"] == 5,
          f"top={d.get('top_wrong_words')}")
    check("U1 平均用时非空",
          (d.get("avg_answer_ms") or 0) > 0 and (d.get("avg_session_ms") or 0) > 0,
          f"answer={d.get('avg_answer_ms')} session={d.get('avg_session_ms')}")
    check("U1 完成率 = 3/4（4 个词里练过 3 个，第 4 个只有空结果）",
          d.get("practiced_word_count") == 3 and d.get("scope_word_count") == 4,
          f"{d.get('practiced_word_count')}/{d.get('scope_word_count')}")

    # ---- U2 单元掌握度 ----
    status, body = _http("GET", f"/api/stats/unit/{uid}", token=admin_token)
    d2 = body.get("data") or {}
    check("U2 单元掌握度 code=0", body.get("code") == 0, f"{str(body)[:160]}")
    check("U2 词条数 = 4", d2.get("word_total") == 4, f"{d2.get('word_total')}")
    check("U2 学生数 = 5", d2.get("student_count") == 5, f"{d2.get('student_count')}")
    check("U2 未完成学生 = 5（都只练了 3/4）",
          len(d2.get("unfinished_students") or []) == 5
          and all(s["missing_words"] == 1 for s in d2["unfinished_students"]),
          f"n={len(d2.get('unfinished_students') or [])}")
    hard = d2.get("hard_words") or []
    check("U2 难词榜按正确率升序",
          [w["accuracy"] for w in hard] == sorted(w["accuracy"] for w in hard),
          f"accs={[w['accuracy'] for w in hard]}")
    check("U2 只有空结果的词不进难词榜",
          seed["word_ids"][3] not in [w["word_id"] for w in hard],
          f"word_ids={[w['word_id'] for w in hard]}")
    check("U2 已掌握比例 = 1/4", d2.get("mastered_word_count") == 1,
          f"mastered={d2.get('mastered_word_count')} ratio={d2.get('mastered_ratio')}")

    # ---- U3 个人报告 ----
    sid = seed["student_ids"][0]
    status, body = _http("GET", f"/api/stats/student/{sid}", token=admin_token)
    d3 = body.get("data") or {}
    check("U3 个人报告 code=0", body.get("code") == 0, f"{str(body)[:160]}")
    check("U3 正确率与班级/单元一致（U5 口径唯一）",
          d3.get("accuracy") == d.get("accuracy") == d2.get("accuracy") == 0.6667,
          f"student={d3.get('accuracy')} class={d.get('accuracy')} unit={d2.get('accuracy')}")
    check("U3 场次 / 练习天数 / 已掌握",
          d3.get("session_count") == 1 and d3.get("practice_days") == 1
          and d3.get("mastered_words") == 1,
          f"session={d3.get('session_count')} days={d3.get('practice_days')} "
          f"mastered={d3.get('mastered_words')}")
    check("U3 最近场次带单元名",
          (d3.get("recent_sessions") or [{}])[0].get("unit_name") == "烟测 Unit 1",
          f"{(d3.get('recent_sessions') or [{}])[0].get('unit_name')}")

    # ---- U4 正确率口径 ----
    check("U4 2 exact + 1 wrong + 1 empty → 2/3（不是 2/4）",
          d3.get("accuracy") == 0.6667 and d3.get("counted_count") == 3,
          f"accuracy={d3.get('accuracy')} counted={d3.get('counted_count')} "
          f"answer={d3.get('answer_count')}")

    # ---- U6 申诉 pending ----
    status, body = _http("GET", "/api/appeals", token=admin_token)
    d6 = body.get("data") or {}
    check("U6 默认只返回待复核", body.get("code") == 0 and d6.get("total") == 1,
          f"total={d6.get('total')}")
    item = (d6.get("items") or [{}])[0]
    check("U6 申诉记录带学生与词条信息",
          item.get("username") and item.get("expected") and item.get("match_type") == "wrong",
          f"username={item.get('username')} match_type={item.get('match_type')}")

    # ---- U8 audio_available ----
    check("U8 audio_available 恒为 false", item.get("audio_available") is False,
          f"{item.get('audio_available')}")

    # ---- U7 复核 ----
    appeal_id = item.get("appeal_id")
    before_is_correct = item.get("is_correct")
    status, body = _http("POST", f"/api/appeals/{appeal_id}/review",
                         token=admin_token,
                         json_body={"result": "upheld", "remark": "烟测：确认学生读对了"})
    check("U7 复核 code=0", body.get("code") == 0, f"{str(body)[:160]}")
    check("U7 返回 is_correct 未变",
          (body.get("data") or {}).get("is_correct") == before_is_correct,
          f"before={before_is_correct} after={(body.get('data') or {}).get('is_correct')}")

    # 换一条连接重新读（证明真的落盘了，不是内存里改了一下）
    status, body = _http("GET", "/api/appeals?status=upheld", token=admin_token)
    dup = body.get("data") or {}
    upheld_item = (dup.get("items") or [{}])[0]
    check("U7 复核结论已落盘（status=upheld 能查到）",
          dup.get("total") == 1 and upheld_item.get("appeal_id") == appeal_id,
          f"total={dup.get('total')}")
    check("U7 ⛔ is_correct 未被改动",
          upheld_item.get("is_correct") == before_is_correct,
          f"before={before_is_correct} now={upheld_item.get('is_correct')}")
    check("U7 复核备注从审计回读到列表",
          upheld_item.get("review_remark") == "烟测：确认学生读对了",
          f"remark={upheld_item.get('review_remark')}")
    check("U7 复核人姓名可见", bool(upheld_item.get("reviewer_name")),
          f"{upheld_item.get('reviewer_name')}")

    status, body = _http("GET", "/api/appeals", token=admin_token)
    check("U7 复核后从待复核列表消失",
          (body.get("data") or {}).get("total") == 0,
          f"total={(body.get('data') or {}).get('total')}")

    # ---- U12 审计查询 ----
    status, body = _http("GET", "/api/audit-logs?page=1&page_size=5", token=admin_token)
    d12 = body.get("data") or {}
    check("U12 审计日志可查", body.get("code") == 0 and d12.get("total") >= 1,
          f"total={d12.get('total')}")
    status, body = _http("GET", "/api/audit-logs?action=appeal.review", token=admin_token)
    d12b = body.get("data") or {}
    check("U12 按动作筛选（前缀匹配）",
          d12b.get("total") == 1
          and d12b["items"][0]["action"] == "appeal.review"
          and (d12b["items"][0]["detail"] or {}).get("remark") == "烟测：确认学生读对了",
          f"total={d12b.get('total')} detail={d12b.get('items', [{}])[0].get('detail')}")
    status, body = _http("GET", "/api/audit-logs?action=%25", token=admin_token)
    check("U12 LIKE 通配符被转义（`%` 不匹配全表）",
          (body.get("data") or {}).get("total") == 0,
          f"total={(body.get('data') or {}).get('total')}")

    # ---- U10 ASR 用量 ----
    status, body = _http("GET", "/api/stats/asr", token=admin_token)
    d10 = body.get("data") or {}
    check("U10 ASR 用量 code=0", body.get("code") == 0, f"{str(body)[:160]}")
    rows = [r for r in (d10.get("daily") or []) if r["provider"] == "aliyun"]
    check("U10 按天 × 厂商有数据", bool(rows), f"daily={d10.get('daily')}")
    if rows:
        r0 = rows[0]
        check("U10 调用次数 = 4（3 ok + 1 error）", r0["call_count"] == 4,
              f"call_count={r0['call_count']}")
        check("U10 失败率 = 1/4（empty 不算失败）", r0["failure_rate"] == 0.25,
              f"failure_rate={r0['failure_rate']}")
        check("U10 平均延迟只算 ok 行 = 250", r0["avg_latency_ms"] == 250,
              f"avg_latency_ms={r0['avg_latency_ms']}")
        check("U10 音频总时长 = 3×1200 + 900", r0["total_audio_ms"] == 3 * 1200 + 900,
              f"total_audio_ms={r0['total_audio_ms']}")
    check("U10 按厂商汇总有 aliyun",
          {p["provider"] for p in (d10.get("by_provider") or [])} == {"aliyun"},
          f"{d10.get('by_provider')}")

    # ---- U13 导出 xlsx ----
    status, content, headers = _http(
        "GET", f"/api/stats/export?type=class&id={cid}", token=admin_token, raw=True
    )
    check("U13 班级导出返回 xlsx", status == 200 and content[:2] == b"PK",
          f"status={status} magic={content[:2]!r} len={len(content)}")
    check("U13 响应头是下载语义",
          "spreadsheetml" in headers.get("content-type", "")
          and "attachment" in headers.get("content-disposition", ""),
          f"{headers.get('content-disposition')}")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        header_names = [c.value for c in ws[1]]
        check("U13 xlsx 能被解析且行数正确（5 学生 + 表头）",
              ws.max_row == 6 and "正确率" in header_names,
              f"max_row={ws.max_row} headers={header_names}")
        idx = {name: i for i, name in enumerate(header_names)}
        first = [c.value for c in ws[2]]
        check("U13 导出的正确率与接口口径一致（66.7%）",
              first[idx["正确率"]] == "66.7%", f"{first[idx['正确率']]}")
    except ImportError:      # pragma: no cover
        check("U13 xlsx 能被解析", False, "openpyxl 不可用")

    status, content, headers = _http(
        "GET", f"/api/stats/export?type=student&id={sid}", token=admin_token, raw=True
    )
    check("U13 学生导出返回 xlsx", status == 200 and content[:2] == b"PK",
          f"status={status} magic={content[:2]!r}")

    # ---- 鉴权 ----
    for path in (f"/api/stats/class/{cid}", "/api/appeals", "/api/audit-logs",
                 "/api/stats/asr"):
        status, body = _http("GET", path)
        check(f"鉴权 无 token {path} → 10002", body.get("code") == 10002,
              f"code={body.get('code')}")
    if student_token:
        status, body = _http("GET", "/api/appeals", token=student_token)
        check("鉴权 学生 token → 10003", body.get("code") == 10003, f"code={body.get('code')}")

    # ---- 参数校验（错误要带字段名）----
    status, body = _http("GET", f"/api/stats/class/{cid}?date_from=20260924",
                         token=admin_token)
    check("参数 紧凑日期格式被拒（YYYYMMDD → 10001 + 字段名）",
          body.get("code") == 10001 and "date_from" in (body.get("message") or ""),
          f"{body.get('code')} {body.get('message')}")
    status, body = _http("GET", f"/api/stats/class/{cid}?top_n=999", token=admin_token)
    check("参数 越界不静默夹值（top_n=999 → 10001）", body.get("code") == 10001,
          f"code={body.get('code')}")
    status, body = _http("GET", f"/api/stats/unit/999999", token=admin_token)
    check("参数 单元不存在 → 30001", body.get("code") == 30001, f"code={body.get('code')}")


def main() -> int:
    if not _tcp_free(PORT):
        print(f"TCP {PORT} 被占用，请用 SMOKE_BE07_PORT 指定其它端口。", flush=True)
        return 2

    # ---- 第一步：造数（独立进程，用应用自己的 DB 层）----
    seed_file = TMP / "seed.py"
    seed_file.write_text(SEED_SCRIPT, encoding="utf-8")
    seeded = subprocess.run(
        [str(PYTHON), str(seed_file)],
        cwd=str(BACKEND), env=ENV, capture_output=True, text=True, encoding="utf-8",
    )
    if seeded.returncode != 0:
        print("造数失败：", flush=True)
        print((seeded.stdout or "")[-1500:], flush=True)
        print((seeded.stderr or "")[-3000:], flush=True)
        shutil.rmtree(TMP, ignore_errors=True)
        return 2
    try:
        seed = json.loads((seeded.stdout or "").strip().splitlines()[-1])
    except (ValueError, IndexError):
        print("造数脚本没有输出预期的 JSON：", flush=True)
        print((seeded.stdout or "")[-1500:], flush=True)
        print((seeded.stderr or "")[-3000:], flush=True)
        shutil.rmtree(TMP, ignore_errors=True)
        return 2

    # ---- 第二步：起真服务 ----
    log_file = TMP / "server.log"
    with log_file.open("w", encoding="utf-8") as lf:
        proc = subprocess.Popen(
            [str(PYTHON), "-m", "uvicorn", "app.main:app",
             "--host", "127.0.0.1", "--port", str(PORT), "--log-level", "warning"],
            cwd=str(BACKEND), env=ENV, stdout=lf, stderr=subprocess.STDOUT,
        )

    try:
        import httpx

        ready = False
        for _ in range(120):
            if proc.poll() is not None:
                break
            try:
                with httpx.Client(timeout=2, trust_env=False) as c:
                    if c.get(BASE + "/api/health").status_code == 200:
                        ready = True
                        break
            except Exception:      # noqa: BLE001
                time.sleep(0.25)
        if not ready:
            print("服务未能就绪，服务端日志：", flush=True)
            print(log_file.read_text(encoding="utf-8", errors="replace")[-3000:], flush=True)
            return 2

        print(f"== BE-07 统计看板烟测（{BASE}）==", flush=True)
        run_checks(seed)
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:      # pragma: no cover
            proc.kill()
        print("== 已收尸（子进程 + 临时目录）==", flush=True)
        shutil.rmtree(TMP, ignore_errors=True)

    print(f"\n烟测结果：{PASS}/{PASS + FAIL} 通过", flush=True)
    if FAILED:
        print("失败项：" + "、".join(FAILED), flush=True)
    print(f"SUMMARY pass={PASS} total={PASS + FAIL} failed={FAIL}", flush=True)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
